"""File text for the model: one converter, explicit gaps, a fixed size budget."""
import io
import json
import pytest
from tools.file_text import extract, share, json_size, MAX_TEXT_BYTES


def pptx(*titles) -> bytes:
    from pptx import Presentation
    deck = Presentation()
    for title in titles:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f'{title} detail'
    out = io.BytesIO(); deck.save(out)
    return out.getvalue()


def xlsx() -> bytes:
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(['Element', 'Mass'])
    book.active.append(['Carbon', 12])
    out = io.BytesIO(); book.save(out)
    return out.getvalue()


async def test_powerpoint_keeps_slide_order_and_titles():
    text, problem = await extract(pptx('The Brain', 'Neurons', 'Synapses'),
                                  'Brain.pptx', 'application/vnd.openxmlformats-officedocument.presentationml.presentation')
    assert problem is None
    assert text.index('The Brain') < text.index('Neurons') < text.index('Synapses')
    assert 'Slide number: 1' in text and 'Neurons detail' in text


async def test_spreadsheet_becomes_a_table():
    text, problem = await extract(xlsx(), 'Masses.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    assert problem is None and '| Carbon | 12 |' in text


async def test_plain_text_and_csv():
    assert (await extract(b'Revise chapter 4', 'notes.txt', 'text/plain')) == ('Revise chapter 4', None)
    text, _ = await extract(b'a,b\n1,2\n', 'data.csv', 'text/csv')
    assert '| a | b |' in text


@pytest.mark.parametrize('data,name,mime,code', [
    (b'\x89PNG\r\n\x1a\n', 'photo.png', 'image/png', 'no_text_layer'),
    (b'PK\x03\x04' + bytes(range(256)) * 8, 'broken.docx',
     'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'conversion_failed'),
])
async def test_files_without_text_say_so(data, name, mime, code):
    assert await extract(data, name, mime) == ('', code)


def test_budget_keeps_short_files_whole_and_splits_the_rest():
    short, long_a, long_b = 'x' * 100, 'é' * 200_000, 'y\n' * 200_000
    fitted = share([short, long_a, long_b])
    assert fitted[0] == short
    assert long_a.startswith(fitted[1]) and long_b.startswith(fitted[2])
    assert sum(json_size(t) for t in fitted) <= MAX_TEXT_BYTES
    assert json_size(fitted[1]) > MAX_TEXT_BYTES // 3 and json_size(fitted[2]) > MAX_TEXT_BYTES // 3
    assert share(['a', 'b']) == ['a', 'b']


async def test_truncated_text_is_flagged_with_the_full_length():
    import httpx
    from test_open_files import school, run, ids_for, pdf
    long_text = 'line of notes\n' * 20_000
    transport = school({'/uploads/asset/file/501/lab.pdf': httpx.Response(200, content=long_text.encode(),
                                                                          headers={'content-type': 'text/plain'})})
    ids = await ids_for(transport)
    result = await run({'class_id': '10', 'task_id': '101', 'open': [ids['Lab instructions.pdf']]}, transport)
    [opened] = result['files']
    assert opened['status'] == 'read' and opened['truncated'] is True
    assert opened['text_chars'] == len(long_text.strip()) and len(opened['text']) < opened['text_chars']
    assert len(json.dumps(result).encode()) < 250_000


async def test_files_after_the_time_budget_are_skipped_not_timed_out(monkeypatch):
    from test_open_files import school, run, ids_for
    import tools.files
    monkeypatch.setattr(tools.files, 'CONVERT_SECONDS', -1)
    ids = await ids_for(school())
    result = await run({'class_id': '10', 'task_id': '101', 'open': [ids['Lab instructions.pdf']]}, school())
    assert result['files'][0]['error']['code'] == 'time_budget'
